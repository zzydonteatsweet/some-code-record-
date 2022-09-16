import openpyxl as op
import re
import pandas as pd
if __name__ == '__main__':
    # ecl = pd.read_excel('附件.xlsx', sheet_name='表单2')
    # ecl = ecl.fillna(0)
    # ecl.to_excel(r'数据2.xlsx', index= False)
    workbook = op.load_workbook(r'数据2.xlsx')
    write_book = op.Workbook()
    wsheet = write_book.active
    sheet = workbook.sheets[0]

    nrow = sheet.max_row ; ncol = sheet.max_column

    for i in range(1, ncol):
        wsheet.cell(1, i).value = sheet.cell(1, i).value
    p = re.compile(r'\S*部位1\S*')
    p2 = re.compile(r'\S*部位2\S*')
    for i in range(2, nrow+1):
        t = str(sheet.cell(i, 1).value)
        if p2.match(t):
            continue
        if p.match(t):
            sheet.cell(i, 1).value = i-1
            for k in range(2, ncol+1):
                tmp = 0
                tmp += float(sheet.cell(i, k).value) ; tmp += float(sheet.cell(i+1, k).value)
                tmp /= 2
                wsheet.cell(i, k).value = tmp
        else:
            for k in range(2, ncol + 1):
                tmp = 0
                tmp += float(sheet.cell(i, k).value);
                wsheet.cell(i, k).value = tmp
    write_book.save('new数据2.xlsx')
